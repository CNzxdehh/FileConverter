# -*- coding: utf-8 -*-
"""自检脚本：生成样本 → 逐条跑通所有转换 → 验证多选合并 → 验证自定义输出目录 → 验证注册表读写。

用法：
    python src\\selftest.py
"""

from __future__ import annotations

import os
import shutil as _sh
import subprocess
import sys
import time
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "src"
SAMPLES = ROOT / "samples"
CUSTOM = SAMPLES / "custom"
CONVERTER = SRC / "converter.py"
CHILD = SRC / "_collect_child.py"
PY = sys.executable

sys.path.insert(0, str(SRC))

# 静默模式：子进程不弹消息框，否则测试会一直等用户点"确定"
os.environ["FILECONV_SILENT"] = "1"

# 设 FC_EXE=dist\\FileConverter.exe 可以直接对打包好的 exe 跑同一套用例
EXE = os.environ.get("FC_EXE")
if EXE:
    CLI = [EXE]
    os.environ["FILECONV_NO_CONSOLE"] = "1"   # exe 是窗口子系统，测试时不弹黑窗口
else:
    CLI = [PY, str(CONVERTER)]

FAILED: list[str] = []
PASSED: list[str] = []


# --------------------------------------------------------------- 工具
def case(name: str, ok: bool, detail: str = "") -> None:
    print(f"  [{'PASS' if ok else 'FAIL'}] {name}" + (f" -> {detail}" if detail else ""))
    (PASSED if ok else FAILED).append(name)


def new_files(dirs, ext: str, before: set[Path]) -> list[Path]:
    after = set()
    for d in dirs:
        if d.exists():
            after |= {p for p in d.iterdir() if p.suffix.lower() == ext.lower()}
    return sorted(after - before)


def snapshot() -> set[Path]:
    out = set()
    for d in (SAMPLES, CUSTOM):
        if d.exists():
            out |= set(d.iterdir())
    return out


def run_cli(args: list[str], timeout: int = 300) -> subprocess.CompletedProcess:
    try:
        return subprocess.run(args, capture_output=True, text=True,
                              encoding="utf-8", errors="replace", timeout=timeout)
    except subprocess.TimeoutExpired:
        print(f"  !! 超时({timeout}s): {' '.join(args)[:100]}")
        return subprocess.CompletedProcess(args, 124, "", "timeout")


def ffmpeg_bin() -> str:
    import shutil

    return shutil.which("ffmpeg") or "ffmpeg"


# --------------------------------------------------------------- 样本
def make_samples() -> dict[str, Path]:
    SAMPLES.mkdir(parents=True, exist_ok=True)
    out: dict[str, Path] = {}

    from PIL import Image, ImageDraw

    png = SAMPLES / "sample.png"
    im = Image.new("RGB", (400, 300), (30, 90, 200))
    ImageDraw.Draw(im).rectangle([40, 40, 360, 260], outline=(255, 255, 255), width=8)
    im.save(png)
    out["png"] = png

    jpg = SAMPLES / "sample.jpg"
    Image.new("RGB", (320, 240), (200, 120, 30)).save(jpg, quality=90)
    out["jpg"] = jpg

    webp = SAMPLES / "sample.webp"
    Image.new("RGB", (256, 256), (60, 180, 90)).save(webp, "WEBP")
    out["webp"] = webp

    pdf = SAMPLES / "sample.pdf"
    import fitz

    for idx, name in ((1, "sample.pdf"), (2, "sample2.pdf"), (3, "sample3.pdf")):
        doc = fitz.open()
        for i in range(1, 3):
            page = doc.new_page(width=595, height=842)
            page.insert_text((72, 100), f"FileConverter self test {name} page {i}", fontsize=22)
            page.insert_text((72, 160), "Hello 你好 123 ABCDEF", fontsize=16)
        doc.save(SAMPLES / name)
        doc.close()
        out[f"pdf{idx}" if idx > 1 else "pdf"] = SAMPLES / name

    docx = SAMPLES / "sample.docx"
    from docx import Document

    d = Document()
    d.add_heading("FileConverter 自检文档", level=1)
    d.add_paragraph("这是一段用于自测的中文文本：转换文件格式工具。")
    d.add_paragraph("Second paragraph for Word to PDF test.")
    d.save(docx)
    out["docx"] = docx

    ff = ffmpeg_bin()
    wav = SAMPLES / "sample.wav"
    run_cli([ff, "-y", "-loglevel", "error", "-f", "lavfi", "-i",
             "sine=frequency=440:duration=3", "-ar", "44100", str(wav)])
    out["wav"] = wav
    for ext, extra in (("flac", []), ("m4a", ["-c:a", "aac", "-b:a", "128k"]),
                       ("ogg", ["-c:a", "libvorbis", "-b:a", "128k"])):
        p = SAMPLES / f"sample.{ext}"
        run_cli([ff, "-y", "-loglevel", "error", "-i", str(wav)] + extra + [str(p)])
        out[ext] = p

    base = [ff, "-y", "-loglevel", "error",
            "-f", "lavfi", "-i", "testsrc=size=320x240:rate=15:duration=3",
            "-f", "lavfi", "-i", "sine=frequency=440:duration=3",
            "-shortest", "-pix_fmt", "yuv420p"]
    for ext in ("mov", "mkv", "avi", "flv"):
        p = SAMPLES / f"sample.{ext}"
        run_cli(base + [str(p)])
        out[ext] = p
    return out


# --------------------------------------------------------------- 用例
def test_collector_unit() -> None:
    print("\n== 多选合并器（并发选举）==")
    n = 6
    probes = [SAMPLES / f"probe{i}.png" for i in range(n)]
    for p in probes:
        p.write_bytes(b"")

    def launch(p: Path) -> str:
        r = run_cli([PY, str(CHILD), str(p), "1200"])
        return (r.stdout or r.stderr or "").strip()

    with ThreadPoolExecutor(max_workers=n) as ex:
        outs = list(ex.map(launch, probes))

    leaders = [o for o in outs if o.startswith("LEADER")]
    counted = [int(o.split()[1]) for o in leaders]
    case(f"{n} 个并发实例只选出 1 个主进程", len(leaders) == 1, str(outs))
    case("主进程收集到全部文件", bool(counted) and counted[0] == n, str(counted))


def test_single(samples: dict[str, Path]) -> None:
    print("\n== 单文件转换 ==")
    table = [
        ("pdf2docx", "pdf", ".docx"),
        ("pdf2txt", "pdf", ".txt"),
        ("docx2pdf", "docx", ".pdf"),
        ("docx2txt", "docx", ".txt"),
        ("pdf2img", "pdf", ".png"),
        ("img2png", "webp", ".png"),
        ("img2jpg", "png", ".jpg"),
        ("img2webp", "jpg", ".webp"),
        ("img2pdf", "png", ".pdf"),
        ("imgcompress", "png", ".png"),
        ("aud2mp3", "wav", ".mp3"),
        ("aud2wav", "m4a", ".wav"),
        ("vid2mp3", "mov", ".mp3"),
        ("vid2mp4", "mov", ".mp4"),
        ("vid2gif", "mov", ".gif"),
        ("aud2mp3", "flac", ".mp3"),
        ("aud2mp3", "m4a", ".mp3"),
        ("aud2mp3", "ogg", ".mp3"),
        ("vid2mp4", "mkv", ".mp4"),
        ("vid2mp4", "avi", ".mp4"),
        ("vid2mp4", "flv", ".mp4"),
    ]
    for verb, key, ext in table:
        before = snapshot()
        src = samples[key]
        r = run_cli([*CLI, "run", verb, str(src)])
        got = new_files([SAMPLES, CUSTOM], ext, before)
        # pdf2img 是"一个 PDF 出多张图"，样本 2 页 → 应得 2 个
        expect = 2 if verb == "pdf2img" else 1
        ok = len(got) == expect and all(p.stat().st_size > 0 for p in got) and r.returncode == 0
        detail = ", ".join(p.name for p in got) if got else f"rc={r.returncode} 无输出"
        case(f"{verb}  ← {src.name}", ok, detail)
        if not ok and r.stdout:
            print("        stdout:", (r.stdout or "")[-400:].replace("\n", " | "))


def test_batch(samples: dict[str, Path]) -> None:
    print("\n== 批量（同时选中多个文件）==")
    before_mp3 = {p for p in SAMPLES.iterdir() if p.suffix == ".mp3"}
    aud = [samples["wav"], samples["flac"], samples["m4a"], samples["ogg"]]
    with ThreadPoolExecutor(max_workers=len(aud)) as ex:
        list(ex.map(lambda f: run_cli([*CLI, "run", "aud2mp3", str(f)]), aud))
    gained = {p for p in SAMPLES.iterdir() if p.suffix == ".mp3"} - before_mp3
    case("多选 4 个音频 → mp3", len(gained) == 4, f"{len(gained)} 个输出")

    before_mp4 = {p for p in SAMPLES.iterdir() if p.suffix == ".mp4"}
    vids = [samples["mkv"], samples["avi"], samples["flv"]]
    with ThreadPoolExecutor(max_workers=3) as ex:
        list(ex.map(lambda f: run_cli([*CLI, "run", "vid2mp4", str(f)]), vids))
    gained = {p for p in SAMPLES.iterdir() if p.suffix == ".mp4"} - before_mp4
    case("多选 3 个视频 → mp4", len(gained) == 3, f"{len(gained)} 个输出")

    before_merge = {p for p in SAMPLES.iterdir() if p.name.startswith("合并_")}
    imgs = [samples["png"], samples["jpg"], samples["webp"]]
    with ThreadPoolExecutor(max_workers=3) as ex:
        list(ex.map(lambda f: run_cli([*CLI, "run", "imgs2pdf", str(f)]), imgs))
    gained = {p for p in SAMPLES.iterdir() if p.name.startswith("合并_")} - before_merge
    ok = len(gained) == 1 and next(iter(gained)).stat().st_size > 1000
    case("多图合并为一个 PDF", ok, f"{len(gained)} 个")

    # 命令行直传多个文件：converter run img2png a.png b.jpg
    before = snapshot()
    r = run_cli([*CLI, "run", "img2png", str(samples["jpg"]), str(samples["webp"])])
    got = new_files([SAMPLES, CUSTOM], ".png", before)
    case("命令行直传 2 个文件", len(got) == 2, str([p.name for p in got]))

    # runm：一次性传入全部文件（一级菜单/COM 走的路径，不做进程选举）
    before = snapshot()
    run_cli([*CLI, "runm", "img2webp", str(samples["png"]), str(samples["jpg"])])
    got = new_files([SAMPLES, CUSTOM], ".webp", before)
    case("runm 一次传 2 个文件", len(got) == 2, str([p.name for p in got]))

    # 多 PDF 合并：单进程直传（一级菜单/runm 走的就是这条路，最可靠）
    before_merge = {p for p in SAMPLES.iterdir() if p.name.startswith("合并_")}
    pdfs = [samples["pdf"], samples["pdf2"], samples["pdf3"]]
    run_cli([*CLI, "runm", "pdfmerge", *[str(p) for p in pdfs]])
    gained = {p for p in SAMPLES.iterdir() if p.name.startswith("合并_")} - before_merge
    pages = 0
    if gained:
        import fitz

        with fitz.open(sorted(gained)[0]) as d:
            pages = d.page_count
    case("多 PDF 合并（runm，页数应为 6）", len(gained) == 1 and pages == 6,
         f"{len(gained)} 个文件 / {pages} 页")

    # 多 PDF 合并：模拟多选（3 个进程）→ 至少要合并出文件，且不丢页
    before_merge = {p for p in SAMPLES.iterdir() if p.name.startswith("合并_")}
    with ThreadPoolExecutor(max_workers=3) as ex:
        list(ex.map(lambda f: run_cli([*CLI, "run", "pdfmerge", str(f)]), pdfs))
    gained = {p for p in SAMPLES.iterdir() if p.name.startswith("合并_")} - before_merge
    total_pages = 0
    if gained:
        import fitz

        for p in gained:
            with fitz.open(p) as d:
                total_pages += d.page_count
    case("多选 PDF 合并（页数合计应为 6）", total_pages == 6,
         f"{len(gained)} 个文件 / 共 {total_pages} 页")


def test_custom_out(samples: dict[str, Path]) -> None:
    print("\n== 自定义输出目录 / 恢复原目录 ==")
    from core.config import set_out_dir

    CUSTOM.mkdir(exist_ok=True)
    src = SAMPLES / "outtest.webp"
    _sh.copyfile(samples["webp"], src)

    set_out_dir(CUSTOM)
    before = snapshot()
    run_cli([*CLI, "run", "img2png", str(src)])
    got = [p for p in snapshot() - before if p.parent == CUSTOM]
    case("文件落在自定义文件夹", len(got) == 1 and got[0].suffix == ".png",
         str([p.name for p in got]))

    set_out_dir(None)
    run_cli([*CLI, "run", "img2jpg", str(src)])
    case("恢复后落在原文件同目录", (SAMPLES / "outtest_转换后.jpg").exists())


def menu_visible() -> bool | None:
    """用 Explorer 自己的动词枚举，确认菜单项真的会出现（None = 环境不支持该检查）。"""
    try:
        import win32com.client
    except Exception:
        return None
    probe = SAMPLES / "probe.txt"
    probe.parent.mkdir(parents=True, exist_ok=True)
    probe.write_text("probe", encoding="utf-8")
    try:
        sh = win32com.client.Dispatch("Shell.Application")
        folder = sh.Namespace(str(SAMPLES))
        item = folder.ParseName(probe.name)
        for v in item.Verbs():
            if "转换文件格式" in (v.Name or ""):
                return True
        return False
    except Exception:
        return None


def test_registry_dry() -> None:
    print("\n== 注册表 / 级联子菜单外壳扩展 ==")
    from core import registry

    was_installed = registry.is_installed()

    run_cli([*CLI, "install"])

    case("install → 级联扩展 CLSID 已注册", registry.ctx_registered())
    ok, msg = registry.ctx_verify()
    case("install → 扩展可加载、调用返回码合法", ok, str(msg))
    case("install → 已挂到文件(*)和文件夹(Directory)", registry.ctx_handlers_bound())
    case("install → 已切到经典完整菜单（级联项才在一级可见）",
         registry.classic_menu_on())

    # CLSID 的 InprocServer32 必须指向 FileConvCtx.dll，且线程模型为 Apartment
    srv = rf"Software\Classes\CLSID\{registry.CLSID_CTX}\InprocServer32"
    dll_val, model = "", ""
    try:
        import winreg

        with winreg.OpenKey(winreg.HKEY_CURRENT_USER, srv) as h:
            dll_val = winreg.QueryValueEx(h, "")[0]
            model = winreg.QueryValueEx(h, "ThreadingModel")[0]
    except OSError:
        pass
    case("InprocServer32 指向 FileConvCtx.dll",
         dll_val.lower().endswith("fileconvctx.dll"), dll_val)
    case("ThreadingModel = Apartment（Explorer 是 STA）", model == "Apartment", model)

    # 两个右键对象的 ContextMenuHandlers 默认值都要等于我们的 CLSID
    bound = True
    for t in registry.CTX_TARGETS:
        key = rf"Software\Classes\{t}\shellex\ContextMenuHandlers\{registry.CTX_HANDLER}"
        try:
            import winreg

            with winreg.OpenKey(winreg.HKEY_CURRENT_USER, key) as h:
                bound = bound and \
                    winreg.QueryValueEx(h, "")[0].upper() == registry.CLSID_CTX.upper()
        except OSError:
            bound = False
    case("ContextMenuHandlers 默认值 = 我们的 CLSID", bound)

    # 旧的 IExplorerCommand / MSIX / 静态项不能被留下来打架
    case("install → 旧 IExplorerCommand 扩展已清除", not registry.com_registered())
    case("install → 旧静态菜单项已清除", not registry.key_exists(registry.PARENT_KEY))

    run_cli([*CLI, "uninstall"])
    case("uninstall → 级联扩展已注销", not registry.ctx_registered())
    case("uninstall → ContextMenuHandlers 已摘掉", not registry.ctx_handlers_bound())
    case("uninstall → is_installed() 为假", not registry.is_installed())
    case("uninstall → 菜单样式已还原", not registry.classic_menu_on())

    if was_installed:      # 别把用户已经装好的菜单测没了
        run_cli([*CLI, "install"])
        print("  [还原] 测试前菜单已安装，已重新装回")


def main(argv: list[str]) -> int:
    t0 = time.time()
    print("== 生成样本文件 ==")
    samples = make_samples()
    print("  ", ", ".join(sorted({p.name for p in samples.values()})))

    test_collector_unit()
    test_single(samples)
    test_batch(samples)
    test_custom_out(samples)
    test_registry_dry()

    print("\n" + "=" * 60)
    print(f"通过 {len(PASSED)}   失败 {len(FAILED)}   用时 {time.time() - t0:.1f}s")
    for f in FAILED:
        print("  失败项:", f)
    print("=" * 60)
    return 1 if FAILED else 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
