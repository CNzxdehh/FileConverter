# -*- coding: utf-8 -*-
"""转换能力矩阵与具体实现。

每个转换由 (源拓展名集合 -> 目标拓展名) 描述，菜单动词一一对应。
"""

from __future__ import annotations

import os
import re
import shutil
import subprocess
import tempfile
import time
import warnings
from pathlib import Path

# PyMuPDF 对 `import fitz` 报 FutureWarning，与用户无关，屏蔽掉保持输出干净
warnings.filterwarnings("ignore", category=FutureWarning, message=".*fitz.*")

from . import ffmpegtool, winui
from .config import Config, out_root, unique_path

CREATE_NO_WINDOW = getattr(subprocess, "CREATE_NO_WINDOW", 0x08000000)

IMG_EXTS = {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif", ".tiff"}
AUD_EXTS = {".wav", ".flac", ".m4a", ".ogg", ".aac", ".wma", ".opus", ".aiff", ".mp3"}
VID_EXTS = {".mov", ".mkv", ".avi", ".flv", ".wmv", ".m4v", ".mpg", ".mpeg", ".ts", ".webm"}

IMG_OUT = {".png": "PNG", ".jpg": "JPEG", ".jpeg": "JPEG", ".webp": "WEBP"}


class ConvertError(RuntimeError):
    pass


# ==========================================================================
# 文档
# ==========================================================================
def _pdf_searchable(src: Path) -> bool:
    try:
        import fitz  # PyMuPDF

        with fitz.open(src) as doc:
            return any(page.get_text().strip() for page in doc)
    except Exception:
        return True


def pdf_to_docx(src: Path, dst: Path, cfg: Config, ff: dict) -> None:
    import logging

    from pdf2docx import Converter  # 会打印大量调试日志

    logging.getLogger("pdf2docx").setLevel(logging.WARNING)
    cv = Converter(str(src))
    try:
        cv.convert(str(dst), start=0, end=None)
    finally:
        cv.close()
    if not dst.exists() or dst.stat().st_size == 0:
        raise ConvertError("转换结果为空（可能是纯扫描件 PDF，建议先做 OCR）")


def pdf_to_txt(src: Path, dst: Path, cfg: Config, ff: dict) -> None:
    import fitz

    with fitz.open(src) as doc:
        if doc.needs_pass:
            raise ConvertError("PDF 已加密，需要密码")
        pages = []
        for i, page in enumerate(doc, 1):
            winui.progress(f"提取文本 {i}/{len(doc)} 页")
            pages.append(page.get_text())
        winui.end_progress()
        text = "\n".join(pages)
    if not text.strip():
        raise ConvertError("未提取到文字，可能是纯扫描件 PDF（建议先 OCR）")
    dst.write_text(text, encoding="utf-8")


def _decode(data: bytes) -> str:
    """子进程输出可能是 UTF-8 也可能是 GBK，两种都试一遍。"""
    if not data:
        return ""
    for enc in ("utf-8", "gbk"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _word_com_to_pdf(src: Path, dst: Path) -> None:
    """通过 Word COM 导出 PDF（需要本机装过 Microsoft Word）。"""
    ps_path = str(src).replace("'", "''")
    ps_out = str(dst).replace("'", "''")
    ps = (
        "$w = $null; $d = $null;"
        "try {"
        "  $w = New-Object -ComObject Word.Application;"
        "  $w.Visible = $false; $w.DisplayAlerts = 0;"
        "  $d = $w.Documents.Open('" + ps_path + "', $false, $true);"
        "  $d.ExportAsFixedFormat('" + ps_out + "', 17);"
        "} catch { write-error $_.Exception.Message; exit 3 }"
        "finally { if ($d) { $d.Close($false) }; if ($w) { $w.Quit() } }"
    )
    r = subprocess.run(
        ["powershell", "-NoProfile", "-STA", "-ExecutionPolicy", "Bypass", "-Command", ps],
        capture_output=True, creationflags=CREATE_NO_WINDOW, timeout=300,
    )
    if r.returncode != 0 or not dst.exists():
        raise ConvertError("Word COM 导出失败：" + _decode(r.stderr).strip()[:200]
                           + "（若没装完整版 Word，会自动改用 LibreOffice / Edge 打印）")


def _libreoffice_to_pdf(src: Path, dst: Path) -> None:
    soffice = shutil.which("soffice")
    if not soffice:
        for p in (r"C:\Program Files\LibreOffice\program\soffice.exe",
                  r"C:\Program Files (x86)\LibreOffice\program\soffice.exe"):
            if Path(p).is_file():
                soffice = p
                break
    if not soffice:
        raise ConvertError("未检测到 Word 或 LibreOffice，无法把 Word 转成 PDF")
    tmpdir = dst.parent / f".soffice_tmp_{os.getpid()}"
    tmpdir.mkdir(exist_ok=True)
    try:
        subprocess.run([soffice, "--headless", "--norestore", "--convert-to", "pdf",
                        "--outdir", str(tmpdir), str(src)],
                       capture_output=True, creationflags=CREATE_NO_WINDOW, timeout=300)
        got = tmpdir / (src.stem + ".pdf")
        if not got.exists():
            raise ConvertError("LibreOffice 未生成 PDF")
        shutil.move(str(got), str(dst))
    finally:
        shutil.rmtree(tmpdir, ignore_errors=True)


def _find_browser() -> Path | None:
    """找 Edge / Chrome，用它的无头模式把 HTML 打印成 PDF。"""
    import winreg

    found: list[Path] = []
    for key in (r"SOFTWARE\Microsoft\Windows\CurrentVersion\App Paths\msedge.exe",
                r"SOFTWARE\Microsoft\Windows\CurrentVersion\App Paths\chrome.exe"):
        try:
            with winreg.OpenKey(winreg.HKEY_LOCAL_MACHINE, key) as h:
                val = winreg.QueryValue(h, None)
                if val:
                    found.append(Path(val.strip().strip('"')))
        except OSError:
            pass
    found += [
        Path(r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe"),
        Path(r"C:\Program Files\Microsoft\Edge\Application\msedge.exe"),
        Path(r"C:\Program Files\Google\Chrome\Application\chrome.exe"),
        Path(r"C:\Program Files (x86)\Google\Chrome\Application\chrome.exe"),
    ]
    for p in found:
        try:
            if p.is_file():
                return p
        except Exception:
            continue
    return None


def _run_html(run) -> str:
    text = (run.text or "").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    if not text.strip():
        return ""
    text = text.replace("\n", "<br>")
    if run.bold:
        text = f"<b>{text}</b>"
    if run.italic:
        text = f"<i>{text}</i>"
    if run.underline:
        text = f"<u>{text}</u>"
    return text


def _paragraph_html(p) -> str:
    inner = "".join(_run_html(r) for r in p.runs)
    if not inner.strip():
        return '<p style="margin:0 0 6pt">&nbsp;</p>'
    style = (p.style.name if p.style is not None else "") or ""
    m = re.match(r"Heading\s*(\d)", style)
    if m:
        lvl = min(int(m.group(1)), 6)
        return f"<h{lvl}>{inner}</h{lvl}>"
    if style.lower().startswith("title"):
        return f"<h1>{inner}</h1>"
    align = getattr(p, "alignment", None)
    css = ""
    if align is not None and str(align).startswith("CENTER"):
        css = ' style="text-align:center"'
    elif align is not None and str(align).startswith("RIGHT"):
        css = ' style="text-align:right"'
    return f"<p{css}>{inner}</p>"


def _table_html(t) -> str:
    rows = []
    for row in t.rows:
        cells = "".join(
            f"<td>{''.join(_paragraph_html(p) for p in c.paragraphs)}</td>"
            for c in row.cells
        )
        rows.append(f"<tr>{cells}</tr>")
    return ('<table border="1" cellspacing="0" cellpadding="5" '
            'style="border-collapse:collapse">{rows}</table>').format(rows="".join(rows))


_HTML_TPL = """<!DOCTYPE html><html><head><meta charset="utf-8">
<style>
  @page {{ size: A4; margin: 20mm 18mm; }}
  body {{ font-family: "Microsoft YaHei", "Segoe UI", sans-serif; font-size: 11pt;
        line-height: 1.6; color: #222; }}
  h1,h2,h3,h4,h5,h6 {{ margin: 12pt 0 6pt; }}
  p {{ margin: 0 0 6pt; }}
  table {{ width: 100%; font-size: 10pt; }}
</style></head><body>{body}</body></html>"""


def _docx_to_pdf_by_print(src: Path, dst: Path) -> None:
    """最后兜底：docx → HTML → Edge/Chrome 无头打印成 PDF（不需要装 Office）。"""
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = Document(str(src))
    parts: list[str] = []
    for child in doc.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            parts.append(_paragraph_html(Paragraph(child, doc)))
        elif tag == "tbl":
            parts.append(_table_html(Table(child, doc)))
    html = _HTML_TPL.format(body="\n".join(parts) or "<p>（空文档）</p>")

    browser = _find_browser()
    if not browser:
        raise ConvertError("未找到 Edge 或 Chrome，无法用打印方式生成 PDF")

    tmp_html = Path(tempfile.gettempdir()) / f"fcprint_{os.getpid()}.html"
    tmp_html.write_text(html, encoding="utf-8")
    profile = Path(tempfile.gettempdir()) / "FileConvPrint"   # 固定复用，避免每次产生垃圾
    profile.mkdir(parents=True, exist_ok=True)
    try:
        if dst.exists():
            dst.unlink()
        args = [str(browser), "--headless=new", "--disable-gpu", "--no-pdf-header-footer",
                "--no-sandbox", "--disable-extensions",
                f"--user-data-dir={profile}",
                "--virtual-time-budget=4000",
                f"--print-to-pdf={dst}",
                tmp_html.resolve().as_uri()]
        subprocess.run(args, capture_output=True, creationflags=CREATE_NO_WINDOW, timeout=180)
        for _ in range(60):  # 等 PDF 落盘
            if dst.exists() and dst.stat().st_size > 0:
                break
            time.sleep(0.5)
        if not dst.exists():
            raise ConvertError("浏览器打印未生成 PDF（可尝试安装 Word 或 LibreOffice）")
    finally:
        try:
            tmp_html.unlink()
        except Exception:
            pass


def docx_to_txt(src: Path, dst: Path, cfg: Config, ff: dict) -> None:
    """Word 转 TXT：正文段落 + 表格文本（按行 tab 分隔）。"""
    from docx import Document

    doc = Document(str(src))
    lines = [p.text.strip() for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            lines.append("\t".join(c.text.strip() for c in row.cells))
    text = "\n".join(x for x in lines if x)
    if not text.strip():
        raise ConvertError("文档里没有可提取的文字")
    dst.write_text(text, encoding="utf-8")


def pdf_merge(srcs: list[Path], dst: Path, cfg: Config, ff: dict) -> None:
    """多个 PDF 合并成一个（保持原有页序）。"""
    import fitz

    out = fitz.open()
    try:
        for i, p in enumerate(srcs, 1):
            winui.progress(f"合并 {i}/{len(srcs)}  {p.name[:30]}")
            with fitz.open(p) as doc:
                if doc.needs_pass:
                    raise ConvertError(f"{p.name} 已加密，需要密码")
                out.insert_pdf(doc)
        winui.end_progress()
        if out.page_count == 0:
            raise ConvertError("没有可合并的页面")
        out.save(str(dst))
    finally:
        out.close()


def pdf_to_images(src: Path, dst: Path, cfg: Config, ff: dict) -> list[Path]:
    """PDF 逐页导出 PNG（150 dpi）。多页时文件名带页号。"""
    import fitz

    outs: list[Path] = []
    with fitz.open(src) as doc:
        n = doc.page_count
        if n == 0:
            raise ConvertError("PDF 没有页面")
        for i, page in enumerate(doc, 1):
            winui.progress(f"渲染第 {i}/{n} 页")
            pix = page.get_pixmap(dpi=150)
            name = f"{dst.stem}_{i}.png" if n > 1 else f"{dst.stem}.png"
            path = unique_path(dst.parent / name)
            pix.save(str(path))
            outs.append(path)
        winui.end_progress()
    return outs


def image_compress(src: Path, dst: Path, cfg: Config, ff: dict) -> None:
    """保持原格式压缩（JPEG/WEBP 质量 80，PNG 走 optimize）。"""
    from PIL import Image, ImageFile

    ImageFile.LOAD_TRUNCATED_IMAGES = True
    fmt = {".png": "PNG", ".jpg": "JPEG", ".jpeg": "JPEG", ".webp": "WEBP"}.get(
        src.suffix.lower()
    )
    if not fmt:
        raise ConvertError("该图片格式不支持压缩（支持 png/jpg/webp）")
    with Image.open(src) as im:
        im.load()
        if fmt == "JPEG" and im.mode not in ("RGB", "L"):
            im = im.convert("RGB")
        kw: dict = {"optimize": True}
        if fmt in ("JPEG", "WEBP"):
            kw["quality"] = 80
        im.save(dst, fmt, **kw)


def audio_to_wav(src: Path, dst: Path, cfg: Config, ff: dict) -> None:
    ok, err = ffmpegtool.run(
        [str(ff["ffmpeg"]), "-i", str(src), "-vn", "-c:a", "pcm_s16le", str(dst)],
        ff.get("total"), f"{src.name[:28]} → wav",
    )
    if not ok:
        raise ConvertError(err)


def video_to_mp3(src: Path, dst: Path, cfg: Config, ff: dict) -> None:
    """从视频里抽音轨转 MP3。"""
    ok, err = ffmpegtool.run(
        [str(ff["ffmpeg"]), "-i", str(src), "-vn", "-map_metadata", "0",
         "-c:a", "libmp3lame", "-b:a", str(cfg.get("mp3_bitrate", "192k")), str(dst)],
        ff.get("total"), f"{src.name[:28]} → mp3",
    )
    if not ok:
        if "does not contain any stream" in err or "matches no streams" in err:
            raise ConvertError("这个视频没有音轨，抽不出音频")
        raise ConvertError(err)


def video_to_gif(src: Path, dst: Path, cfg: Config, ff: dict) -> None:
    """视频转 GIF：12fps、宽 480、调色板优化。"""
    vf = ("fps=12,scale=480:-1:flags=lanczos,split[s0][s1];"
          "[s0]palettegen=max_colors=200[p];[s1][p]paletteuse=dither=bayer")
    ok, err = ffmpegtool.run(
        [str(ff["ffmpeg"]), "-i", str(src), "-vf", vf, "-loop", "0", str(dst)],
        ff.get("total"), f"{src.name[:28]} → gif",
    )
    if not ok:
        raise ConvertError(err)


def word_installed() -> bool:
    try:
        import winreg

        winreg.CloseKey(winreg.OpenKey(winreg.HKEY_CLASSES_ROOT, "Word.Application"))
        return True
    except Exception:
        return False


def docx_to_pdf(src: Path, dst: Path, cfg: Config, ff: dict) -> None:
    """三级方案：Word COM → LibreOffice → Edge/Chrome 打印（都不行才报错）。"""
    if src.suffix.lower() == ".doc":
        raise ConvertError("旧版 .doc 请先另存为 .docx（.doc 无法直接解析）")

    errors: list[str] = []
    if word_installed():
        try:
            _word_com_to_pdf(src, dst)
            return
        except Exception as e:
            errors.append(str(e)[:160])
    else:
        errors.append("未检测到 Microsoft Word")

    try:
        _libreoffice_to_pdf(src, dst)
        return
    except Exception as e:
        errors.append(str(e)[:160])

    try:
        _docx_to_pdf_by_print(src, dst)
        return
    except Exception as e:
        errors.append(str(e)[:160])

    raise ConvertError("Word 转 PDF 失败，三种方案都不可用：\n  " + "\n  ".join(errors))


# ==========================================================================
# 图片
# ==========================================================================
def image_convert(src: Path, dst: Path, fmt_key: str, cfg: Config) -> None:
    from PIL import Image, ImageFile

    ImageFile.LOAD_TRUNCATED_IMAGES = True
    with Image.open(src) as im:
        im.load()
        target = IMG_OUT[fmt_key]
        if target == "JPEG":
            if im.mode in ("RGBA", "LA", "P"):
                bg = Image.new("RGB", im.size, (255, 255, 255))
                base = im.convert("RGBA")
                bg.paste(base, mask=base.split()[-1])
                im = bg
            else:
                im = im.convert("RGB")
        try:
            exif = im.getexif()
        except Exception:
            exif = None
        save_kw = {}
        if target == "JPEG":
            save_kw.update(quality=95, subsampling=0, optimize=True)
        elif target == "WEBP":
            save_kw.update(quality=92, method=4)
        if exif and target in ("JPEG", "WEBP"):
            try:
                save_kw["exif"] = exif
            except Exception:
                pass
        im.save(dst, target, **save_kw)


def image_to_pdf(src: Path, dst: Path, cfg: Config, ff: dict) -> None:
    from PIL import Image, ImageFile

    ImageFile.LOAD_TRUNCATED_IMAGES = True
    with Image.open(src) as im:
        im.load()
        rgb = Image.new("RGB", im.size, (255, 255, 255))
        if im.mode == "RGBA":
            rgb.paste(im.convert("RGBA"), mask=im.convert("RGBA").split()[-1])
        else:
            rgb.paste(im.convert("RGB"))
        rgb.save(dst, "PDF", resolution=150.0)


def images_to_pdf(srcs: list[Path], dst: Path, cfg: Config, ff: dict) -> None:
    from PIL import Image, ImageFile

    ImageFile.LOAD_TRUNCATED_IMAGES = True
    opened = []
    try:
        for i, p in enumerate(srcs, 1):
            winui.progress(f"合并图片 {i}/{len(srcs)}")
            im = Image.open(p)
            im.load()
            opened.append(im.convert("RGB"))
        winui.end_progress()
        opened[0].save(dst, "PDF", save_all=True, append_images=opened[1:],
                       resolution=150.0)
    finally:
        for im in opened:
            try:
                im.close()
            except Exception:
                pass


# ==========================================================================
# 音视频
# ==========================================================================
def audio_to_mp3(src: Path, dst: Path, cfg: Config, ff: dict) -> None:
    ffmpeg = ff["ffmpeg"]
    ok, err = ffmpegtool.run(
        [str(ffmpeg), "-i", str(src), "-vn", "-map_metadata", "0",
         "-c:a", "libmp3lame", "-b:a", str(cfg.get("mp3_bitrate", "192k")), str(dst)],
        ff.get("total"), f"{src.name[:28]} → mp3",
    )
    if not ok:
        raise ConvertError(err)


def video_to_mp4(src: Path, dst: Path, cfg: Config, ff: dict) -> None:
    ffmpeg = ff["ffmpeg"]
    # H.264 + AAC，yuv420p 保证播放器兼容
    ok, err = ffmpegtool.run(
        [str(ffmpeg), "-i", str(src),
         "-map", "0:v:0?", "-map", "0:a?", "-sn", "-dn",
         "-c:v", "libx264", "-preset", "veryfast", "-crf", str(cfg.get("mp4_crf", "23")),
         "-pix_fmt", "yuv420p", "-vf", "scale=trunc(iw/2)*2:trunc(ih/2)*2",
         "-c:a", "aac", "-b:a", "192k", "-movflags", "+faststart", str(dst)],
        ff.get("total"), f"{src.name[:28]} → mp4",
    )
    if not ok:
        raise ConvertError(err)


def probe_duration(src: Path, ffmpeg: Path) -> None:
    return None  # 占位，实际调用见 runner


# ==========================================================================
# 动词表：菜单项 = 这里的一条
# ==========================================================================
# accept   : 接受的源拓展名（小写集合）
# out_ext  : 目标拓展名
# handler  : 单文件处理函数
# needs_ff : 是否需要 ffmpeg
SPECS: dict[str, dict] = {
    "pdf2docx": dict(
        label="PDF 转 Word (.docx)", accept={".pdf"}, out_ext=".docx",
        handler=pdf_to_docx, needs_ff=False, group="文档",
    ),
    "docx2pdf": dict(
        label="Word 转 PDF", accept={".docx", ".doc"}, out_ext=".pdf",
        handler=docx_to_pdf, needs_ff=False, group="文档",
    ),
    "pdf2txt": dict(
        label="PDF 转 TXT", accept={".pdf"}, out_ext=".txt",
        handler=pdf_to_txt, needs_ff=False, group="文档",
    ),
    "docx2txt": dict(
        label="Word 转 TXT", accept={".docx"}, out_ext=".txt",
        handler=docx_to_txt, needs_ff=False, group="文档",
    ),
    "pdf2img": dict(
        label="PDF 逐页转 PNG", accept={".pdf"}, out_ext=".png",
        handler=pdf_to_images, needs_ff=False, group="文档", multi_out=True,
    ),
    "pdfmerge": dict(
        label="多PDF合并 → 一个 PDF", accept={".pdf"}, out_ext=".pdf",
        handler=None, batch_handler=pdf_merge, needs_ff=False, group="文档",
        batch_kind="个 PDF",
    ),
    "img2png": dict(
        label="图片 → PNG", accept=IMG_EXTS, out_ext=".png",
        handler=lambda s, d, c, f: image_convert(s, d, ".png", c),
        needs_ff=False, group="图片",
    ),
    "img2jpg": dict(
        label="图片 → JPG", accept=IMG_EXTS, out_ext=".jpg",
        handler=lambda s, d, c, f: image_convert(s, d, ".jpg", c),
        needs_ff=False, group="图片",
    ),
    "img2webp": dict(
        label="图片 → WEBP", accept=IMG_EXTS, out_ext=".webp",
        handler=lambda s, d, c, f: image_convert(s, d, ".webp", c),
        needs_ff=False, group="图片",
    ),
    "img2pdf": dict(
        label="图片 → PDF", accept=IMG_EXTS, out_ext=".pdf",
        handler=image_to_pdf, needs_ff=False, group="图片",
    ),
    "imgs2pdf": dict(
        label="多图合并 → 一个 PDF", accept=IMG_EXTS, out_ext=".pdf",
        handler=None, batch_handler=images_to_pdf, needs_ff=False, group="图片",
        batch_kind="张图片",
    ),
    "imgcompress": dict(
        label="图片压缩（保原格式）", accept={".png", ".jpg", ".jpeg", ".webp"},
        out_ext="", suffix="_压缩后", handler=image_compress, needs_ff=False, group="图片",
    ),
    "aud2mp3": dict(
        label="音频 → MP3", accept=AUD_EXTS, out_ext=".mp3",
        handler=audio_to_mp3, needs_ff=True, group="音视频",
    ),
    "aud2wav": dict(
        label="音频 → WAV", accept=AUD_EXTS, out_ext=".wav",
        handler=audio_to_wav, needs_ff=True, group="音视频",
    ),
    "vid2mp3": dict(
        label="视频提取音频 → MP3", accept=VID_EXTS, out_ext=".mp3",
        handler=video_to_mp3, needs_ff=True, group="音视频",
    ),
    "vid2mp4": dict(
        label="视频 → MP4", accept=VID_EXTS, out_ext=".mp4",
        handler=video_to_mp4, needs_ff=True, group="音视频",
    ),
    "vid2gif": dict(
        label="视频 → GIF", accept=VID_EXTS, out_ext=".gif",
        handler=video_to_gif, needs_ff=True, group="音视频",
    ),
}

# 额外功能（不参与文件处理，单独走 utilities 分支）
UTILS = {
    "setoutdir": "自定义输出文件夹…",
    "cleardir": "恢复输出到原目录",
    "checkenv": "检查运行环境",
    "ffmpeghelp": "ffmpeg 下载指引",
    "about": "关于本工具",
}

VERB_ORDER = [
    # 文档
    "pdf2docx", "docx2pdf", "pdf2txt", "docx2txt", "pdf2img", "pdfmerge",
    # 图片
    "img2png", "img2jpg", "img2webp", "img2pdf", "imgs2pdf", "imgcompress",
    # 音频 / 视频
    "aud2mp3", "aud2wav", "vid2mp3", "vid2mp4", "vid2gif",
]


def out_name(src: Path, spec: dict, cfg: Config) -> Path:
    """输出路径：原名 + 后缀 + 目标拓展名。

    spec 里可以覆盖后缀（例如压缩用 _压缩后），out_ext 为空表示保持原拓展名。
    """
    suffix = spec.get("suffix") or cfg.get("suffix", "_转换后")
    ext = spec.get("out_ext") or src.suffix
    return out_root(cfg, src) / f"{src.stem}{suffix}{ext}"


def build_ff(cfg: Config, src: Path | None = None) -> dict:
    ffmpeg = ffmpegtool.find_ffmpeg(cfg.get("ffmpeg_path"))
    ff = {"ffmpeg": ffmpeg, "ffprobe": ffmpegtool.find_ffprobe(ffmpeg) if ffmpeg else None}
    if src and ff["ffprobe"]:
        try:
            ff["total"] = ffmpegtool.duration(ff["ffprobe"], src)
        except Exception:
            ff["total"] = None
    return ff


def convert_file(verb: str, src: Path, cfg: Config, ffmpeg: Path | None = None
                 ) -> tuple[Path, str | None, list[Path]]:
    """返回 (主输出路径, 错误信息, 全部输出文件)。错误信息为 None 表示成功。

    ffmpeg 可传入已定位好的路径做缓存，避免每个文件重复搜索。
    有些转换会产出多个文件（PDF 逐页转 PNG），全部输出放在第三个返回值里。
    """
    spec = SPECS[verb]
    dst = unique_path(out_name(src, spec, cfg))
    dst.parent.mkdir(parents=True, exist_ok=True)
    ff: dict = {}
    if spec.get("needs_ff"):
        ffmpeg = ffmpeg or ffmpegtool.find_ffmpeg(cfg.get("ffmpeg_path"))
        ff = {"ffmpeg": ffmpeg,
              "ffprobe": ffmpegtool.find_ffprobe(ffmpeg) if ffmpeg else None}
        if ff["ffprobe"]:
            try:
                ff["total"] = ffmpegtool.duration(ff["ffprobe"], src)
            except Exception:
                ff["total"] = None
    try:
        res = spec["handler"](src, dst, cfg, ff)
        outs = [p for p in (res if isinstance(res, list) else [dst])
                if isinstance(p, Path) and p.exists() and p.stat().st_size > 0]
        if not outs:
            return dst, "未生成有效输出文件", []
        outs.sort()
        return outs[0], None, outs
    except Exception as e:
        try:
            if dst.exists():
                dst.unlink()
        except Exception:
            pass
        return dst, f"{type(e).__name__}: {e}"[:400], []
